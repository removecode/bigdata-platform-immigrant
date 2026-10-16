from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT = Path("大数据平台YARN上的任务梳理.xlsx")
OUTPUT_DIR = Path("outputs")
OUTPUT = OUTPUT_DIR / "项目数据来源及输出_全集join.xlsx"

NS = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def column_index(cell_ref: str) -> int:
    letters = "".join(re.findall(r"[A-Z]+", cell_ref))
    idx = 0
    for char in letters:
        idx = idx * 26 + ord(char) - 64
    return idx - 1


def load_shared_strings(zf: ZipFile) -> list[str]:
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    values: list[str] = []
    for item in root.findall("a:si", NS):
        values.append("".join(text.text or "" for text in item.findall(".//a:t", NS)))
    return values


def cell_text(cell: ET.Element, shared_strings: list[str]) -> str:
    cell_type = cell.attrib.get("t")
    value = cell.find("a:v", NS)
    if cell_type == "s":
        return shared_strings[int(value.text)] if value is not None and value.text else ""
    if cell_type == "inlineStr":
        return "".join(text.text or "" for text in cell.findall(".//a:t", NS))
    return value.text if value is not None and value.text else ""


def read_sheet(zf: ZipFile, sheet_path: str) -> list[list[str]]:
    shared_strings = load_shared_strings(zf)
    root = ET.fromstring(zf.read(f"xl/{sheet_path}"))
    rows: list[list[str]] = []
    max_cols = 0
    for row in root.findall(".//a:sheetData/a:row", NS):
        values: list[str] = []
        for cell in row.findall("a:c", NS):
            idx = column_index(cell.attrib["r"])
            while len(values) <= idx:
                values.append("")
            values[idx] = cell_text(cell, shared_strings)
            max_cols = max(max_cols, idx + 1)
        rows.append(values)
    for row in rows:
        row.extend([""] * (max_cols - len(row)))
    return rows


def unique_headers(headers: list[str], prefix: str) -> list[str]:
    counts: Counter[str] = Counter()
    result: list[str] = []
    for idx, header in enumerate(headers, start=1):
        name = str(header).strip() or f"未命名列{idx}"
        counts[name] += 1
        if counts[name] > 1:
            name = f"{name}_{counts[name]}"
        result.append(f"{prefix}{name}")
    return result


def trim_blank_columns(rows: list[list[str]]) -> list[list[str]]:
    max_used = 0
    for row in rows:
        for idx, value in enumerate(row):
            if str(value).strip():
                max_used = max(max_used, idx + 1)
    return [row[:max_used] for row in rows]


def normalize_key(value: object) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^[*＊\\s]+", "", text)
    text = re.sub(r"\\s+", " ", text)
    return text.casefold()


def rows_to_dicts(rows: list[list[str]], header_row: int, data_start: int, prefix: str) -> list[dict[str, str]]:
    rows = trim_blank_columns(rows)
    headers = unique_headers(rows[header_row], prefix)
    dicts: list[dict[str, str]] = []
    for raw in rows[data_start:]:
        raw = raw + [""] * (len(headers) - len(raw))
        if not any(str(value).strip() for value in raw):
            continue
        dicts.append({headers[idx]: raw[idx] for idx in range(len(headers))})
    return dicts


def autofit(ws, max_width: int = 70) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        width = 10
        for row_idx in range(1, min(ws.max_row, 200) + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is None:
                continue
            lines = str(value).splitlines() or [""]
            width = max(width, min(max_width, max(len(line) for line in lines) + 2))
        ws.column_dimensions[letter].width = width


def write_table(ws, headers: list[str], rows: list[list[object]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit(ws)


def main() -> None:
    with ZipFile(INPUT) as zf:
        sheet1_rows = read_sheet(zf, "worksheets/sheet1.xml")
        sheet3_rows = read_sheet(zf, "worksheets/sheet3.xml")

    sheet1 = rows_to_dicts(sheet1_rows, header_row=1, data_start=2, prefix="云脑_")
    sheet3 = rows_to_dicts(sheet3_rows, header_row=0, data_start=1, prefix="来源_")

    service_col = "云脑_服务名称"
    project_col = "来源_项目名称"

    current_service = ""
    for row in sheet1:
        raw_service = str(row.get(service_col, "")).strip()
        if raw_service:
            current_service = raw_service
        row["云脑_服务名称_填充"] = current_service
        row["join_key"] = normalize_key(current_service)

    for row in sheet3:
        row["join_key"] = normalize_key(row.get(project_col, ""))

    left_by_key: dict[str, list[dict[str, str]]] = defaultdict(list)
    right_by_key: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in sheet1:
        if row["join_key"]:
            left_by_key[row["join_key"]].append(row)
    for row in sheet3:
        if row["join_key"]:
            right_by_key[row["join_key"]].append(row)

    all_keys = sorted(set(left_by_key) | set(right_by_key))
    left_cols = ["云脑_服务名称_填充"] + [col for col in sheet1[0].keys() if col not in {"join_key", "云脑_服务名称_填充"}]
    right_cols = [col for col in sheet3[0].keys() if col != "join_key"]
    headers = ["join_key", "join状态"] + left_cols + right_cols

    joined_rows: list[list[object]] = []
    status_counter: Counter[str] = Counter()
    for key in all_keys:
        left_rows = left_by_key.get(key, [])
        right_rows = right_by_key.get(key, [])
        if left_rows and right_rows:
            status = "两边匹配"
            pairs = [(left, right) for left in left_rows for right in right_rows]
        elif left_rows:
            status = "仅sheet1"
            pairs = [(left, {}) for left in left_rows]
        else:
            status = "仅sheet3"
            pairs = [({}, right) for right in right_rows]
        status_counter[status] += len(pairs)
        for left, right in pairs:
            joined_rows.append([key, status] + [left.get(col, "") for col in left_cols] + [right.get(col, "") for col in right_cols])

    wb = Workbook()
    ws = wb.active
    ws.title = "全集数据"
    write_table(ws, headers, joined_rows)

    stats = wb.create_sheet("匹配统计")
    stats_rows = [
        ["输入文件", INPUT.name],
        ["sheet1", "云脑大数据平台"],
        ["sheet3", "项目数据来源及输出"],
        ["join 方式", "full outer join"],
        ["sheet1 join 字段", "服务名称（空白行向下填充后使用）"],
        ["sheet3 join 字段", "项目名称"],
        ["sheet1 原始数据行数", len(sheet1)],
        ["sheet3 原始数据行数", len(sheet3)],
        ["输出全集行数", len(joined_rows)],
        ["两边匹配行数", status_counter["两边匹配"]],
        ["仅sheet1行数", status_counter["仅sheet1"]],
        ["仅sheet3行数", status_counter["仅sheet3"]],
        ["sheet1唯一服务数", len(left_by_key)],
        ["sheet3唯一项目数", len(right_by_key)],
        ["匹配唯一键数", len(set(left_by_key) & set(right_by_key))],
        ["仅sheet1唯一键数", len(set(left_by_key) - set(right_by_key))],
        ["仅sheet3唯一键数", len(set(right_by_key) - set(left_by_key))],
    ]
    write_table(stats, ["指标", "值"], stats_rows)

    unmatched1 = wb.create_sheet("仅sheet1")
    unmatched1_rows = [
        [key, *[row.get(col, "") for col in left_cols]]
        for key in sorted(set(left_by_key) - set(right_by_key))
        for row in left_by_key[key]
    ]
    write_table(unmatched1, ["join_key", *left_cols], unmatched1_rows)

    unmatched3 = wb.create_sheet("仅sheet3")
    unmatched3_rows = [
        [key, *[row.get(col, "") for col in right_cols]]
        for key in sorted(set(right_by_key) - set(left_by_key))
        for row in right_by_key[key]
    ]
    write_table(unmatched3, ["join_key", *right_cols], unmatched3_rows)

    OUTPUT_DIR.mkdir(exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)
    print(dict(status_counter))


if __name__ == "__main__":
    main()
