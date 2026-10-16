from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "需求" / "大数据平台-运营需求-收集.xlsx"
OUT_MD = ROOT / "outputs" / "大数据平台运营需求原始诉求分析临时文档.md"
OUT_XLSX = ROOT / "outputs" / "大数据平台运营需求原始诉求结构化提取.xlsx"


TABLE_RE = re.compile(
    r"\b(?:from|join)\s+([`\"']?)([a-zA-Z_][\w-]*(?:\.[a-zA-Z_][\w-]*)?)\1",
    re.IGNORECASE,
)
INLINE_TABLE_RE = re.compile(r"(?:Hive表|数据表)[:：]?\s*[`'\"]([^`'\"]+\.[^`'\"]+)[`'\"]")
COND_RE = re.compile(
    r"(?im)^\s*(?:where|and|or)\s+(.+?)(?=\s*(?:--.*)?$)"
)
FIELD_RE = re.compile(
    r"(?<![\w.])([a-zA-Z_][\w]*|str\d+|long\d+|day|hour|year|month|oid)\s*"
    r"(?:=|>=|<=|>|<|!=|<>|\bin\b|\bnot\s+like\b|\blike\b|\bregexp_like\b|\brlike\b)",
    re.IGNORECASE,
)
FUNC_FIELD_RE = re.compile(r"\b(?:regexp_like|json_extract_scalar)\s*\(\s*([a-zA-Z_][\w]*)", re.IGNORECASE)
TIME_RE = re.compile(
    r"(?i)\b(?:day|hour|year|month|oid|log_time)\b\s*"
    r"(?:=|>=|<=|>|<|between|in)\s*[^)\n;]+"
)
CAST_HOUR_RE = re.compile(r"(?i)cast\s*\(\s*hour\s+as\s+int\s*\)\s*(?:=|>=|<=|>|<)\s*\d+")
RESERVED_FIELDS = {"and", "or", "where", "from", "select", "group", "order", "limit"}
TABLE_EQUIVALENTS = {
    "default.log_audit_source": "csc.ods_csc_log_audit",
    "csc.ods_csc_log_audit": "csc.ods_csc_log_audit",
    "default.log_stats_source": "csc.ods_csc_log_stats",
    "csc.ods_csc_log_stats": "csc.ods_csc_log_stats",
}
FIELD_VALUE_PATTERN = (
    r"(?is)\b({fields})\b\s*(?:=|in)\s*(\([^)]*\)|'[^']*'|\"[^\"]*\"|[a-zA-Z0-9_.$${{}}:-]+)"
)


def clean_sql(sql: str) -> str:
    sql = str(sql or "")
    # Keep markdown SQL blocks usable, but remove quote markers and headings from analyzers.
    lines = []
    for line in sql.splitlines():
        stripped = line.strip()
        if stripped.startswith(">") or stripped.startswith("|"):
            continue
        if stripped.startswith("```"):
            continue
        lines.append(line)
    return "\n".join(lines)


def strip_sql_comments(sql: str) -> str:
    lines = []
    for line in sql.splitlines():
        if "--" in line:
            line = line.split("--", 1)[0]
        lines.append(line)
    return "\n".join(lines)


def normalize_table(name: str) -> str:
    return name.strip().strip("`'\"").lower()


def extract_tables(sql: str) -> list[str]:
    tables = {normalize_table(m.group(2)) for m in TABLE_RE.finditer(sql)}
    tables.update(normalize_table(m.group(1)) for m in INLINE_TABLE_RE.finditer(sql))
    return sorted(t for t in tables if "." in t)


def extract_conditions(sql: str) -> list[str]:
    conditions = []
    for match in COND_RE.finditer(sql):
        cond = re.sub(r"\s+", " ", match.group(1)).strip()
        cond = re.sub(r"\s+--.*$", "", cond).strip()
        if cond and cond not in conditions:
            conditions.append(cond)
    return conditions


def extract_fields(sql: str, conditions: list[str]) -> list[str]:
    fields = set()
    target = "\n".join(conditions) or sql
    for match in FIELD_RE.finditer(target):
        field = match.group(1).lower()
        if field not in RESERVED_FIELDS:
            fields.add(field)
    for match in FUNC_FIELD_RE.finditer(target):
        field = match.group(1).lower()
        if field not in RESERVED_FIELDS:
            fields.add(field)
    return sorted(fields)


def extract_time_filters(sql: str) -> list[str]:
    filters = []
    for match in TIME_RE.finditer(sql):
        expr = re.sub(r"\s+", " ", match.group(0)).strip()
        expr = re.split(r"(?i)\s+(?:and|group\s+by|order\s+by|limit)\s+", expr)[0].strip()
        if expr not in filters:
            filters.append(expr)
    for match in CAST_HOUR_RE.finditer(sql):
        expr = re.sub(r"\s+", " ", match.group(0)).strip()
        if expr not in filters:
            filters.append(expr)
    return filters


def split_sql_values(raw: str) -> list[str]:
    raw = raw.strip()
    if raw.startswith("(") and raw.endswith(")"):
        raw = raw[1:-1]
    values = []
    for match in re.finditer(r"'([^']*)'|\"([^\"]*)\"|([^,\s]+)", raw):
        value = next((g for g in match.groups() if g), "").strip()
        if value:
            values.append(value)
    return values


def extract_filter_values(sql: str, fields: list[str]) -> dict[str, list[str]]:
    if not fields:
        return {}
    pattern = FIELD_VALUE_PATTERN.format(fields="|".join(re.escape(f) for f in fields))
    found: dict[str, list[str]] = defaultdict(list)
    for match in re.finditer(pattern, sql, re.IGNORECASE | re.DOTALL):
        field = match.group(1).lower()
        for value in split_sql_values(match.group(2)):
            if value not in found[field]:
                found[field].append(value)
    return dict(found)


def canonicalize_tables(tables: list[str]) -> list[str]:
    canonical = {TABLE_EQUIVALENTS.get(table, table) for table in tables}
    return sorted(canonical)


def data_flow_for_tables(tables: list[str], purpose: str, sql: str) -> tuple[list[str], list[str]]:
    canonical = set(canonicalize_tables(tables))
    text = f"{purpose}\n{sql}".lower()
    flows: set[str] = set()
    result_layers: set[str] = set()

    if "csc.ods_csc_log_audit" in canonical:
        flows.add("log_compress：default.log_audit_source -> csc.ods_csc_log_audit")
        result_layers.add("深圳旧日志 ODS 审计中间表")
    if "csc.ods_csc_log_stats" in canonical:
        flows.add("log_compress：default.log_stats_source -> csc.ods_csc_log_stats")
        result_layers.add("深圳旧日志 ODS 统计中间表")
    if any(t in canonical for t in ["default.dwb_md5_info", "default.dwb_url_info", "default.dwb_domain_info"]):
        flows.add("log_compress audit DWB + merge_data.sh：csc.ods_csc_log_audit -> default.dwb_*_info")
        result_layers.add("旧链路主 DWB 情报宽表")
    if any(t.startswith("sodl.dwb_sodl_") for t in canonical):
        flows.add("sodl_ods_info_new -> sodl_dwb_info：SODL ODS -> sodl.dwb_sodl_*_info")
        result_layers.add("SODL 新数据湖 DWB 情报宽表")
    if "sri.ods_sri_hashinfo" in canonical:
        flows.add("威胁情报维表：sri.ods_sri_hashinfo")
        result_layers.add("文件样本/MD5 情报维表")
    if "topthreat" in text or "热度" in text:
        if any(t.startswith("sodl.dwb_sodl_") or t.startswith("default.dwb_") for t in canonical):
            result_layers.add("运营热度查询依赖 DWB 明细结果，不直接查询 topthreat_hot_stat Mongo 输出")

    if not flows:
        flows.add("未能从关键数据流文档映射到明确链路")
    if not result_layers:
        result_layers.add("未明确")
    return sorted(flows), sorted(result_layers)


def classify_topic(purpose: str, sql: str) -> str:
    text = f"{purpose}\n{sql}".lower()
    if any(k in text for k in ["沙箱", "sandbox", "cape", "triage"]):
        return "沙箱/流量送检"
    if any(k in text for k in ["url", "下载", "域名", "domain", "合规"]):
        return "URL/域名/合规过滤"
    if any(k in text for k in ["云查", "上报", "文件", "hash", "md5", "样本"]):
        return "文件/样本/hash 查询"
    if any(k in text for k in ["ip主探", "主探", "ip热度", "ip-realtime", "ip-scheduling"]):
        return "IP 主探/信誉/时延分析"
    if any(k in text for k in ["多引擎", "multi-engine"]):
        return "多引擎处理量统计"
    return "其他运营查询"


def result_layers_for_flow(flow: str) -> list[str]:
    if "default.log_audit_source -> csc.ods_csc_log_audit" in flow:
        return ["深圳旧日志 ODS 审计中间表"]
    if "default.log_stats_source -> csc.ods_csc_log_stats" in flow:
        return ["深圳旧日志 ODS 统计中间表"]
    if "default.dwb_*_info" in flow:
        return ["旧链路主 DWB 情报宽表", "运营热度查询依赖 DWB 明细结果，不直接查询 topthreat_hot_stat Mongo 输出"]
    if "sodl.dwb_sodl_*_info" in flow:
        return ["SODL 新数据湖 DWB 情报宽表", "运营热度查询依赖 DWB 明细结果，不直接查询 topthreat_hot_stat Mongo 输出"]
    if "sri.ods_sri_hashinfo" in flow:
        return ["文件样本/MD5 情报维表"]
    return ["未明确"]


def summarize_need(purpose: str, tables: list[str], fields: list[str], time_filters: list[str]) -> str:
    table_text = "、".join(tables) if tables else "未显式识别表"
    field_text = "、".join(fields[:10]) if fields else "未显式识别过滤字段"
    time_text = "；".join(time_filters[:4]) if time_filters else "依赖表格查询范围或脚本参数"
    return f"{purpose}：查询 {table_text}；关键过滤字段 {field_text}；时间条件 {time_text}。"


def autofit(ws, widths: dict[int, int] | None = None) -> None:
    widths = widths or {}
    for col_idx, cells in enumerate(ws.iter_cols(), start=1):
        max_len = max(len(str(cell.value or "")) for cell in cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(widths.get(col_idx, max_len + 2), 70)


def write_workbook(
    records: list[dict],
    common_rows: list[tuple[str, str]],
    flow_summary_rows: list[tuple[str, int, str, str]],
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "原始诉求结构化"
    headers = [
        "序号",
        "目的",
        "主题分类",
        "原始查询表",
        "归一数据源/等价表",
        "应用ID字段",
        "应用ID筛选值",
        "module_id筛选值",
        "关键过滤字段",
        "SQL时间条件",
        "表格查询范围",
        "查询频率",
        "查询方式",
        "填写人",
        "备注",
        "依赖数据流",
        "依赖结果层",
        "原始诉求摘要",
    ]
    ws.append(headers)
    for r in records:
        ws.append(
            [
                r["idx"],
                r["purpose"],
                r["topic"],
                "\n".join(r["tables"]),
                "\n".join(r["canonical_tables"]),
                "app_id/log_app_id" if r["app_values"] else "",
                "\n".join(r["app_values"]),
                "\n".join(r["module_values"]),
                "、".join(r["fields"]),
                "\n".join(r["time_filters"]),
                r["range"],
                r["frequency"],
                r["method"],
                r["owner"],
                r["remark"],
                "\n".join(r["data_flows"]),
                "\n".join(r["result_layers"]),
                r["summary"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True, name="Arial")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit(ws, {2: 28, 4: 34, 5: 34, 7: 36, 8: 34, 9: 42, 10: 44, 16: 48, 17: 42, 18: 70})

    ws2 = wb.create_sheet("共性诉求")
    ws2.append(["类别", "提炼结果"])
    for item in common_rows:
        ws2.append(list(item))
    for cell in ws2[1]:
        cell.font = Font(bold=True, name="Arial")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.freeze_panes = "A2"
    autofit(ws2, {1: 22, 2: 90})

    ws3 = wb.create_sheet("依赖数据流汇总")
    ws3.append(["依赖数据流", "需求条数", "依赖结果层", "覆盖诉求"])
    for item in flow_summary_rows:
        ws3.append(list(item))
    for cell in ws3[1]:
        cell.font = Font(bold=True, name="Arial")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws3.freeze_panes = "A2"
    autofit(ws3, {1: 58, 2: 12, 3: 42, 4: 90})
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)


def main() -> None:
    df = pd.read_excel(INPUT, dtype=str).fillna("")
    records = []
    table_counter: Counter[str] = Counter()
    field_counter: Counter[str] = Counter()
    topic_counter: Counter[str] = Counter()
    freq_counter: Counter[str] = Counter()
    method_counter: Counter[str] = Counter()
    range_counter: Counter[str] = Counter()
    owner_counter: Counter[str] = Counter()
    canonical_table_counter: Counter[str] = Counter()
    app_value_counter: Counter[str] = Counter()
    module_value_counter: Counter[str] = Counter()
    flow_counter: Counter[str] = Counter()
    table_to_needs: defaultdict[str, list[str]] = defaultdict(list)
    flow_to_needs: defaultdict[str, list[str]] = defaultdict(list)
    flow_to_layers: defaultdict[str, set[str]] = defaultdict(set)

    for i, row in df.iterrows():
        purpose = row["目的"].strip()
        sql = clean_sql(row["查询SQL"])
        active_sql = strip_sql_comments(sql)
        tables = extract_tables(active_sql)
        canonical_tables = canonicalize_tables(tables)
        conditions = extract_conditions(active_sql)
        fields = extract_fields(active_sql, conditions)
        time_filters = extract_time_filters(active_sql)
        filter_values = extract_filter_values(active_sql, ["app_id", "log_app_id", "module_id"])
        app_values = sorted(set(filter_values.get("app_id", []) + filter_values.get("log_app_id", [])))
        module_values = filter_values.get("module_id", [])
        data_flows, result_layers = data_flow_for_tables(tables, purpose, active_sql)
        topic = classify_topic(purpose, sql)
        rec = {
            "idx": i + 1,
            "purpose": purpose,
            "sql": sql,
            "topic": topic,
            "tables": tables,
            "canonical_tables": canonical_tables,
            "fields": fields,
            "app_values": app_values,
            "module_values": module_values,
            "time_filters": time_filters,
            "data_flows": data_flows,
            "result_layers": result_layers,
            "range": row["查询范围（当前往前推）"].strip(),
            "frequency": row["查询频率"].strip(),
            "method": row["查询方式"].strip(),
            "owner": row["填写人"].strip(),
            "remark": row["备注"].strip(),
            "summary": summarize_need(purpose, tables, fields, time_filters),
        }
        records.append(rec)
        table_counter.update(tables)
        canonical_table_counter.update(canonical_tables)
        field_counter.update(fields)
        app_value_counter.update(app_values)
        module_value_counter.update(module_values)
        flow_counter.update(data_flows)
        topic_counter.update([topic])
        freq_counter.update([rec["frequency"]])
        method_counter.update([rec["method"]])
        range_counter.update([rec["range"]])
        owner_counter.update([rec["owner"]])
        for table in tables:
            table_to_needs[table].append(purpose)
        for flow in data_flows:
            flow_to_needs[flow].append(purpose)
            flow_to_layers[flow].update(result_layers_for_flow(flow))

    table_lines = [f"{table}（{count} 条）" for table, count in table_counter.most_common()]
    canonical_table_lines = [f"{table}（{count} 条）" for table, count in canonical_table_counter.most_common()]
    topic_lines = [f"{topic}：{count} 条" for topic, count in topic_counter.most_common()]
    field_lines = [f"{field}（{count} 条）" for field, count in field_counter.most_common(20)]
    app_value_lines = [f"{value}（{count} 条）" for value, count in app_value_counter.most_common()]
    module_value_lines = [f"{value}（{count} 条）" for value, count in module_value_counter.most_common()]
    flow_lines = [f"{flow}（{count} 条）" for flow, count in flow_counter.most_common()]
    freq_lines = [f"{freq}：{count} 条" for freq, count in freq_counter.most_common()]
    method_lines = [f"{method}：{count} 条" for method, count in method_counter.most_common()]
    range_lines = [f"{rng}：{count} 条" for rng, count in range_counter.most_common()]
    owner_lines = [f"{owner}：{count} 条" for owner, count in owner_counter.most_common()]

    common_rows = [
        ("高频查询表", "；".join(table_lines)),
        ("归一数据源/等价表", "；".join(canonical_table_lines)),
        ("主题分布", "；".join(topic_lines)),
        ("高频过滤字段", "；".join(field_lines)),
        ("应用ID筛选值 app_id/log_app_id", "；".join(app_value_lines)),
        ("module_id 筛选值", "；".join(module_value_lines)),
        ("依赖数据流", "；".join(flow_lines)),
        ("查询范围", "；".join(range_lines)),
        ("查询频率", "；".join(freq_lines)),
        ("查询方式", "；".join(method_lines)),
        ("填写人分布", "；".join(owner_lines)),
        (
            "共性诉求 1：统一日志查询入口",
            "大量诉求集中在 default.log_audit_source、csc.ods_csc_log_audit、csc.ods_csc_log_stats、default.log_stats_source，且都围绕 app_id/log_app_id、module_id、day/hour、str*/long* 字段过滤，适合沉淀一个按业务/模块/时间/IOC 条件检索的运营查询入口。",
        ),
        (
            "共性诉求 2：IOC 精准检索",
            "文件 MD5、URL、域名、IP/IP:Port 是主要检索对象，当前散落在 str2/str3/str7/str10、json_extract_scalar、regexp_like/RLIKE 和 DWB 宽表字段中，需要标准化 IOC 类型、值、匹配方式和字段映射。",
        ),
        (
            "共性诉求 3：时间窗口参数化",
            "多数查询按近一月回溯，少数一天或跨 hour 窗口；SQL 中同时存在固定 day、${start_filter_day}/${end_filter_day}、${v1s_day}/${v2e_day}、oid 区间等写法，需要统一 start_day/end_day/hour/oid 范围参数。",
        ),
        (
            "共性诉求 4：统计与明细并存",
            "诉求既包括明细追踪（任务状态、上报日志、送检记录、未检出链路、P95 时延），也包括聚合统计（热度、处理量、TOP Hash、趋势），平台能力应支持明细下钻和 group by 指标输出。",
        ),
        (
            "共性诉求 5：人工运营自动化",
            "全部记录为手工或临时脚本，且多项多次/天、多次/周；主探相关备注明确提到过去高强度人工运营和 2.0 自动化需求，是优先自动化场景。",
        ),
    ]

    flow_summary_rows = [
        (
            flow,
            count,
            "；".join(sorted(flow_to_layers[flow])),
            "；".join(dict.fromkeys(flow_to_needs[flow])),
        )
        for flow, count in flow_counter.most_common()
    ]

    write_workbook(records, common_rows, flow_summary_rows)

    md = []
    md.append("# 大数据平台运营需求原始诉求分析临时文档\n")
    md.append(f"- 来源文件：`{INPUT}`")
    md.append(f"- 工作表：`工作表1`")
    md.append(f"- 需求记录数：{len(records)}")
    md.append(f"- 结构化 Excel：`{OUT_XLSX}`")
    md.append("")
    md.append("## 一、总体统计\n")
    md.append("| 维度 | 结果 |")
    md.append("| --- | --- |")
    md.append(f"| 主题分布 | {'；'.join(topic_lines)} |")
    md.append(f"| 高频查询表 | {'；'.join(table_lines)} |")
    md.append(f"| 归一数据源/等价表 | {'；'.join(canonical_table_lines)} |")
    md.append(f"| 高频过滤字段 Top20 | {'；'.join(field_lines)} |")
    md.append(f"| 应用ID筛选值 app_id/log_app_id | {'；'.join(app_value_lines)} |")
    md.append(f"| module_id 筛选值 | {'；'.join(module_value_lines)} |")
    md.append(f"| 依赖数据流 | {'；'.join(flow_lines)} |")
    md.append(f"| 查询范围 | {'；'.join(range_lines)} |")
    md.append(f"| 查询频率 | {'；'.join(freq_lines)} |")
    md.append(f"| 查询方式 | {'；'.join(method_lines)} |")
    md.append(f"| 填写人 | {'；'.join(owner_lines)} |")
    md.append("")
    md.append("## 二、逐条原始诉求结构化\n")
    md.append("| # | 原始目的 | 主题 | 原始查询表 | 归一数据源 | app_id/log_app_id 筛选值 | module_id 筛选值 | 时间范围/条件 | 依赖数据流 |")
    md.append("| --- | --- | --- | --- | --- | --- | --- | --- | --- |")
    for r in records:
        tables = "<br>".join(r["tables"]) if r["tables"] else "-"
        canonical_tables = "<br>".join(r["canonical_tables"]) if r["canonical_tables"] else "-"
        app_values = "<br>".join(r["app_values"]) if r["app_values"] else "-"
        module_values = "<br>".join(r["module_values"]) if r["module_values"] else "-"
        time_text = "<br>".join(r["time_filters"]) if r["time_filters"] else r["range"] or "-"
        data_flows = "<br>".join(r["data_flows"]) if r["data_flows"] else "-"
        md.append(
            f"| {r['idx']} | {r['purpose']} | {r['topic']} | {tables} | {canonical_tables} | {app_values} | {module_values} | {time_text} | {data_flows} |"
        )
    md.append("")
    md.append("## 三、按表归并的诉求\n")
    for table, purposes in table_to_needs.items():
        md.append(f"- `{table}`：{len(purposes)} 条，" + "；".join(dict.fromkeys(purposes)))
    md.append("")
    md.append("## 四、用户依赖的数据流结果\n")
    for flow, count, layers, purposes in flow_summary_rows:
        md.append(f"- **{flow}**：{count} 条；结果层：{layers}；覆盖诉求：{purposes}")
    md.append("")
    md.append("## 五、共性诉求提炼\n")
    for category, text in common_rows[11:]:
        md.append(f"### {category}")
        md.append(text)
        md.append("")
    md.append("## 六、可落地的平台能力建议\n")
    md.append("1. 建立统一查询表单：业务主题、IOC 类型、IOC 值、start_day/end_day、hour、app_id/log_app_id、module_id、输出模式。")
    md.append("2. 建立字段映射层：将 str1~str10、long1~long5 映射为业务含义，尤其是文件云查、文件上报、IP 主探、URL 下载、沙箱送检场景。")
    md.append("3. 建立可复用模板：任务状态查询、日志明细查询、IOC 热度查询、处理量趋势、链路诊断、P95 时延分析。")
    md.append("4. 优先自动化高频场景：IP 主探过滤/降误报/提交统计、hash/IP/域名热度、多次/月的文件上报与云查日志。")
    md.append("5. 对大范围查询做保护：强制时间范围、限制 TOP N、支持异步导出，避免运营手工 SQL 扫大表。")
    md.append("")
    md.append("## 七、原始 SQL 数据质量观察\n")
    md.append("- 第 13 条 `主探过滤模块提交日志` 中出现 `day = 20251507`，月份疑似非法，后续产品化时应做日期校验。")
    md.append("- 第 21 条 `沙箱处理量` 中 `str4='cape' /'triage'` 不是常规 SQL 条件写法，疑似想表达 `str4 in ('cape','triage')` 或二选一条件。")
    md.append("- 多条查询依赖 `str1~str10`、`long1~long5` 的隐式业务含义，原始表缺少字段语义层会导致运营 SQL 难复用、难校验。")
    OUT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")

    print(f"records={len(records)}")
    print(f"tables={len(table_counter)}")
    print(f"fields={len(field_counter)}")
    print(f"wrote={OUT_MD}")
    print(f"wrote={OUT_XLSX}")


if __name__ == "__main__":
    main()
