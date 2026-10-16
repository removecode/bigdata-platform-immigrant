from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
MATCH_DOC = ROOT / "outputs" / "app-call-graph匹配统计临时文档.md"
APP_RESULTS = Path("/Users/jun/project/sxf/2026/app-call-graph/results")
DETAIL_DIR = APP_RESULTS / "detail"
WIKI_DOCS_DIR = APP_RESULTS / "wiki_docs" / "docs"
OUT_MD = ROOT / "outputs" / "app-call-graph匹配项目服务上下文提取临时文档.md"
OUT_XLSX = ROOT / "outputs" / "app-call-graph匹配项目服务上下文提取.xlsx"

DOC_TYPES = ["需求规格说明书.md", "组件依赖分析.md", "项目交互关系.md"]
TABLE_RE = re.compile(r"`([a-zA-Z_][\w-]*\.[a-zA-Z_][\w-]*)`")
PLAIN_TABLE_RE = re.compile(r"\b([a-zA-Z_][\w-]*\.[a-zA-Z_][\w-]*)\b")
KNOWN_DB_PREFIXES = {
    "default",
    "csc",
    "sodl",
    "sri",
    "ngaf",
    "sip",
    "ods",
    "dwb",
    "scia_analysis",
    "customerassets",
}
NON_TABLE_SUFFIXES = {
    "py",
    "scala",
    "java",
    "go",
    "sh",
    "md",
    "json",
    "yaml",
    "yml",
    "conf",
    "properties",
    "svc",
    "cluster",
    "local",
    "novalocal",
    "com",
    "cn",
    "org",
    "net",
    "exe",
    "jar",
}
NON_TABLE_NAMES = {
    "default.format_json_list",
}


def parse_code_block_after_heading(text: str, heading: str) -> list[str]:
    marker = f"## {heading}"
    start = text.find(marker)
    if start < 0:
        return []
    block_start = text.find("```text", start)
    if block_start < 0:
        block_start = text.find("```", start)
    if block_start < 0:
        return []
    block_end = text.find("```", block_start + 3)
    if block_end < 0:
        return []
    block = text[block_start:block_end].split("\n", 1)[1]
    return [line.strip() for line in block.splitlines() if line.strip()]


def matched_names() -> list[str]:
    text = MATCH_DOC.read_text(encoding="utf-8")
    names = set(parse_code_block_after_heading(text, "已匹配服务名称"))
    names.update(parse_code_block_after_heading(text, "已匹配项目名称"))
    return sorted(names)


def load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def compact(text: str, limit: int = 260) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


def markdown_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {}
    current = "前言"
    sections[current] = []
    for line in text.splitlines():
        if line.startswith("## "):
            current = line.lstrip("#").strip()
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


def first_paragraph_after_heading(text: str, heading: str) -> str:
    idx = text.find(heading)
    if idx < 0:
        return ""
    tail = text[idx + len(heading) :]
    paras = [p.strip() for p in re.split(r"\n\s*\n", tail) if p.strip()]
    for para in paras:
        if para.startswith("|") or para.startswith("---"):
            continue
        return compact(re.sub(r"[*#`]", "", para), 360)
    return ""


def extract_project_overview(doc_text: str) -> str:
    for heading in ["## 模块概述", "## 项目概述"]:
        summary = first_paragraph_after_heading(doc_text, heading)
        if summary:
            return summary
    match = re.search(r"\*\*模块定位\*\*[:：]\s*(.+)", doc_text)
    if match:
        return compact(match.group(1), 360)
    match = re.search(r"\|\s*主要功能\s*\|\s*([^|]+)\|", doc_text)
    if match:
        return compact(match.group(1), 360)
    return ""


def normalize_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(v) for v in value if str(v)]
    if isinstance(value, tuple | set):
        return [str(v) for v in value if str(v)]
    text = str(value).strip()
    if not text or text == "-":
        return []
    return [text]


def extract_rules(doc_text: str) -> str:
    rules = []
    for line in doc_text.splitlines():
        stripped = line.strip()
        if "[规则" in stripped or stripped.startswith("* **[规则"):
            rules.append(re.sub(r"[*`]", "", stripped).strip("- "))
        elif stripped.startswith("### 需求点"):
            rules.append(re.sub(r"#+\s*", "", stripped))
        if len(rules) >= 8:
            break
    return "\n".join(rules)


def extract_business_flow(interaction_text: str, req_text: str) -> str:
    parts = []
    for text, headings in [
        (interaction_text, ["## 1. 上游依赖", "## 2. 下游依赖", "## 项目概述"]),
        (req_text, ["## 业务流程", "## 需求点列表"]),
    ]:
        for heading in headings:
            section = first_paragraph_after_heading(text, heading)
            if section:
                parts.append(f"{heading.replace('## ', '')}：{section}")
            if len(parts) >= 4:
                break
    return "\n".join(parts)


def extract_tables_from_text(text: str) -> list[str]:
    tables = set()
    for candidate in set(TABLE_RE.findall(text)) | set(PLAIN_TABLE_RE.findall(text)):
        normalized = candidate.strip("`'\"")
        parts = normalized.split(".")
        if len(parts) != 2:
            continue
        left, right = parts[0], parts[1]
        left_l = left.lower()
        right_l = right.lower()
        if normalized in NON_TABLE_NAMES:
            continue
        if right.endswith("_") or right_l in {"main"}:
            continue
        if right_l in NON_TABLE_SUFFIXES or left_l in {"self", "config", "logger", "grpc", "os", "json"}:
            continue
        if any(right_l.endswith(f".{suffix}") for suffix in NON_TABLE_SUFFIXES):
            continue
        if left_l in KNOWN_DB_PREFIXES or left_l.startswith("db_") or right_l.startswith(("tb_", "ods_", "dwb_")):
            tables.add(normalized)
    return sorted(tables)


def resource_to_table(resource_name: str) -> str:
    if "@" not in resource_name:
        return ""
    parts = [p for p in resource_name.split("@") if p]
    if len(parts) >= 3:
        return f"{parts[-2]}.{parts[-1]}"
    if len(parts) >= 2:
        return ".".join(parts[-2:])
    return ""


def normalize_table_candidate(candidate: str) -> str:
    if not candidate:
        return ""
    if "." not in candidate:
        return ""
    return extract_tables_from_text(candidate)[0] if extract_tables_from_text(candidate) else ""


def extract_resources(detail: dict[str, Any]) -> tuple[list[str], list[str], list[str]]:
    resources = []
    tables = []
    services = []
    for hit in detail.get("hits") or []:
        rtype = str(hit.get("resource_type", ""))
        rname = str(hit.get("resource_name", ""))
        confidence = str(hit.get("confidence", ""))
        if rname:
            resources.append(f"{rtype}:{rname}({confidence})")
        if rtype in {"mongo", "mysql", "hive", "clickhouse", "table"}:
            table = resource_to_table(rname)
            table = normalize_table_candidate(table)
            if table:
                tables.append(table)
        if rtype in {"svc", "grpc", "http", "service"} and rname:
            services.append(rname)
    return sorted(set(resources)), sorted(set(tables)), sorted(set(services))


def wiki_docs_for(name: str) -> dict[str, str]:
    base = WIKI_DOCS_DIR / name
    docs = {}
    for doc_type in DOC_TYPES:
        p = base / doc_type
        if p.exists():
            docs[doc_type] = p.read_text(encoding="utf-8")
    return docs


def wiki_docs_by_content_match(name: str) -> tuple[str, dict[str, str]]:
    for service_dir in WIKI_DOCS_DIR.iterdir() if WIKI_DOCS_DIR.exists() else []:
        if not service_dir.is_dir():
            continue
        docs = {}
        combined = []
        for doc_type in DOC_TYPES:
            p = service_dir / doc_type
            if p.exists():
                text = p.read_text(encoding="utf-8")
                docs[doc_type] = text
                combined.append(text)
        if docs and name in "\n".join(combined):
            return service_dir.name, docs
    return "", {}


def find_content_matches(name: str) -> list[Path]:
    matches = []
    for root in [APP_RESULTS / "wiki_docs", APP_RESULTS]:
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if path.is_file() and path.suffix in {".md", ".json"}:
                try:
                    if name in path.read_text(encoding="utf-8", errors="ignore"):
                        matches.append(path)
                except UnicodeDecodeError:
                    continue
            if len(matches) >= 8:
                break
        if matches:
            break
    return matches


def autofit(ws) -> None:
    for idx, cells in enumerate(ws.iter_cols(), start=1):
        max_len = max(len(str(c.value or "")) for c in cells)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 12), 70)


def write_xlsx(records: list[dict[str, Any]], flow_rows: list[tuple[str, int, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "项目服务上下文"
    headers = [
        "名称",
        "匹配类型",
        "资料来源",
        "介绍",
        "规范/规则摘要",
        "业务流程摘要",
        "依赖表",
        "资源依赖",
        "服务依赖",
        "日志 app_id",
        "日志 module_id",
        "部署数据中心",
        "语言",
        "备注",
    ]
    ws.append(headers)
    for r in records:
        ws.append(
            [
                r["name"],
                r["match_type"],
                "\n".join(r["sources"]),
                r["overview"],
                r["rules"],
                r["business_flow"],
                "\n".join(r["tables"]),
                "\n".join(r["resources"]),
                "\n".join(r["services"]),
                "\n".join(r["app_ids"]),
                "\n".join(r["module_ids"]),
                "\n".join(r["datacenters"]),
                r["language"],
                r["note"],
            ]
        )
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True, name="Arial")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit(ws)

    ws2 = wb.create_sheet("依赖表汇总")
    ws2.append(["依赖表", "项目/服务数量", "项目/服务"])
    for row in flow_rows:
        ws2.append(list(row))
    for cell in ws2[1]:
        cell.font = Font(bold=True, name="Arial")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.freeze_panes = "A2"
    autofit(ws2)
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)


def main() -> None:
    names = matched_names()
    service_names = set(parse_code_block_after_heading(MATCH_DOC.read_text(encoding="utf-8"), "已匹配服务名称"))
    project_names = set(parse_code_block_after_heading(MATCH_DOC.read_text(encoding="utf-8"), "已匹配项目名称"))
    records = []
    table_to_names: defaultdict[str, set[str]] = defaultdict(set)
    source_counter: Counter[str] = Counter()

    for name in names:
        detail = load_json(DETAIL_DIR / f"{name}.json")
        docs = wiki_docs_for(name)
        doc_alias = ""
        if not docs:
            doc_alias, docs = wiki_docs_by_content_match(name)
        req = docs.get("需求规格说明书.md", "")
        comp = docs.get("组件依赖分析.md", "")
        inter = docs.get("项目交互关系.md", "")
        resources, resource_tables, services = extract_resources(detail)
        doc_tables = extract_tables_from_text("\n".join(docs.values()))
        tables = sorted(set(resource_tables + doc_tables))
        for table in tables:
            table_to_names[table].add(name)

        logs = detail.get("logs_info") or {}
        deploy = detail.get("deploy") or {}
        sources = []
        if detail:
            sources.append(f"detail/{name}.json")
        for doc_type in DOC_TYPES:
            if doc_type in docs:
                sources.append(f"wiki_docs/docs/{doc_alias or name}/{doc_type}")
        if not sources:
            content_matches = find_content_matches(name)
            sources.extend(str(p.relative_to(APP_RESULTS)) for p in content_matches[:5])

        source_counter.update(["有wiki文档" if docs else "无wiki文档"])
        source_counter.update(["有detail资源" if detail else "无detail资源"])
        match_type = []
        if name in service_names:
            match_type.append("服务名称")
        if name in project_names:
            match_type.append("项目名称")

        note = ""
        if not docs:
            note = "未找到 wiki_docs/docs 同名目录，介绍/流程主要来自 detail 资源或内容命中。"
        if name == "create_hive_partition":
            note = "项目名通过内容命中到 create-hive-partition-job 文档；无同名 detail 资源。"

        records.append(
            {
                "name": name,
                "match_type": "、".join(match_type),
                "sources": sources,
                "overview": extract_project_overview(req or inter or comp),
                "rules": extract_rules(req),
                "business_flow": extract_business_flow(inter, req),
                "tables": tables,
                "resources": resources[:40],
                "services": services,
                "app_ids": sorted(set(normalize_list(logs.get("app_id")))),
                "module_ids": sorted(set(normalize_list(logs.get("module_id")))),
                "datacenters": sorted(set(normalize_list(deploy.get("datacenters")))),
                "language": str(detail.get("language") or ""),
                "note": note,
            }
        )

    table_rows = [
        (table, len(names_for_table), "；".join(sorted(names_for_table)))
        for table, names_for_table in sorted(table_to_names.items(), key=lambda item: (-len(item[1]), item[0]))
    ]
    write_xlsx(records, table_rows)

    md = []
    md.append("# app-call-graph 匹配项目/服务上下文提取临时文档\n")
    md.append(f"- 匹配统计来源：`{MATCH_DOC}`")
    md.append(f"- app-call-graph 结果目录：`{APP_RESULTS}`")
    md.append(f"- 去重匹配项目/服务数：{len(records)}")
    md.append(f"- 结构化 Excel：`{OUT_XLSX}`")
    md.append("")
    md.append("## 一、覆盖情况\n")
    md.append("| 维度 | 数量 |")
    md.append("| --- | ---: |")
    md.append(f"| 有 wiki_docs 同名文档 | {sum(1 for r in records if any('wiki_docs/docs/' in s for s in r['sources']))} |")
    md.append(f"| 有 detail 资源扫描 | {sum(1 for r in records if any(s.startswith('detail/') for s in r['sources']))} |")
    md.append(f"| 识别到依赖表的项目/服务 | {sum(1 for r in records if r['tables'])} |")
    md.append("")
    md.append("## 二、逐项目/服务提取\n")
    for r in records:
        md.append(f"### {r['name']}")
        md.append(f"- 匹配类型：{r['match_type'] or '-'}")
        md.append(f"- 资料来源：{'; '.join(r['sources']) if r['sources'] else '-'}")
        md.append(f"- 介绍：{r['overview'] or '-'}")
        md.append(f"- 规范/规则摘要：{compact(r['rules'], 700) or '-'}")
        md.append(f"- 业务流程摘要：{compact(r['business_flow'], 700) or '-'}")
        md.append(f"- 依赖表：{', '.join(r['tables']) if r['tables'] else '-'}")
        md.append(f"- 服务依赖：{', '.join(r['services']) if r['services'] else '-'}")
        md.append(f"- 日志字段：app_id={','.join(r['app_ids']) or '-'}；module_id={','.join(r['module_ids']) or '-'}")
        if r["note"]:
            md.append(f"- 备注：{r['note']}")
        md.append("")
    md.append("## 三、依赖表汇总 Top 50\n")
    md.append("| 依赖表 | 项目/服务数量 | 项目/服务 |")
    md.append("| --- | ---: | --- |")
    for table, count, names_text in table_rows[:50]:
        md.append(f"| `{table}` | {count} | {names_text} |")
    OUT_MD.write_text("\n".join(md) + "\n", encoding="utf-8")

    print(f"records={len(records)}")
    print(f"with_docs={sum(1 for r in records if any('wiki_docs/docs/' in s for s in r['sources']))}")
    print(f"with_detail={sum(1 for r in records if any(s.startswith('detail/') for s in r['sources']))}")
    print(f"with_tables={sum(1 for r in records if r['tables'])}")
    print(f"tables={len(table_rows)}")
    print(f"wrote={OUT_MD}")
    print(f"wrote={OUT_XLSX}")


if __name__ == "__main__":
    main()
