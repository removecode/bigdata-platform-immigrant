#!/usr/bin/env python3
"""生成 outputs/业务SQL跑通记录.xlsx。

初始 SQL 取自 917 梳理 CSV 的「语句」列，只收「业务分类」非空的行。
热门威胁的最终 SQL 取自 outputs/topthreat-hot-stat-starrocks.sql。
文件已存在时拒绝覆盖，避免冲掉人工补录的内容；确需重建时加 --force。
"""
import argparse
import csv
import re
import sys
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "大数据平台YARN上的任务梳理-云脑大数据平台-917梳理.csv"
FINAL_SQL_PATH = ROOT / "outputs" / "topthreat-hot-stat-starrocks.sql"
OUT_PATH = ROOT / "outputs" / "业务SQL跑通记录.xlsx"

HOT = "热门威胁-热度统计"
CELL_LIMIT = 32767

DONE = "已跑通"
TODO = "未记录"

HOT_NOTES = {
    "black_domain_stat": (
        "库表 default.log_audit_source 改 csc.ods_csc_log_audit（DEC-014）；log_app_id 改 app_id（DEC-015）；"
        "sodl 的 day/hour 改 insertdate='2026-10-01' 与 hour(from_unixtime(inserttimestamp))=0（DEC-016/017）；"
        "sum(log_num) 改 count(1)（DEC-020）；MANAGE/manage 改 deviceId（DEC-021）；"
        "LATERAL VIEW explode(format_json_list()) 改 unnest(parse_json())，size()-1 改 cardinality()，数组下标改 1 起（ACT-038）；"
        "加 SET_VAR(new_planner_optimize_timeout=60000)（DEC-022）；${day}/${hour} 写死为实验窗口 20261001 / 0。"
    ),
    "blcak_file_stat": (
        "库表改 csc.ods_csc_log_audit（DEC-014）；log_app_id 改 app_id（DEC-015）；"
        "explode(format_json_list()) 改 unnest(parse_json())（ACT-038）；"
        "加 SET_VAR(new_planner_optimize_timeout=60000)（DEC-022）。"
        "该语句没有 sodl 分支。旧 SQL 写死 day=20250826、hour=01，新语句写 day=20261001、hour=0（实验窗口）。"
    ),
    "blcak_url_stat": (
        "改写项同 black_domain_stat：库表、app_id、sodl 时间条件、count(1)、deviceId、unnest(parse_json())、规划超时提示；"
        "${day}/${hour} 写死为实验窗口 20261001 / 0。"
    ),
}

HOT_EVIDENCE = (
    "用户确认已跑通（2026-10-16）。本机实验库另执行过三条：黑域名 110 行、黑文件 0 行、黑 URL 100 行"
    "（outputs/topthreat-hot-stat-results.tsv，ACT-042）。黑文件 0 行是本机造数的 long1 不为 0，不是语句报错。"
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
DONE_FILL = PatternFill("solid", fgColor="C6EFCE")
TODO_FILL = PatternFill("solid", fgColor="FFF2CC")


def clip(text):
    return text if len(text) <= CELL_LIMIT else text[: CELL_LIMIT - 20] + "\n...（超出单元格上限，已截断）"


def load_initial_rows():
    with open(CSV_PATH, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    header = rows[1]
    idx = {name: header.index(name) for name in ("服务名称", "任务类型", "子任务", "业务类别", "业务分类", "语句", "负责人")}
    service = ""
    out = []
    for r in rows[2:]:
        if r[idx["服务名称"]].strip():
            service = r[idx["服务名称"]].strip().lstrip("*")
        biz = r[idx["业务分类"]].strip()
        if not biz:
            continue
        sql = r[idx["语句"]].strip()
        if sql in ("", "/"):
            sql = ""
        out.append(
            {
                "biz": biz,
                "category": r[idx["业务类别"]].strip(),
                "service": service,
                "status": r[idx["任务类型"]].strip(),
                "owner": r[idx["负责人"]].strip(),
                "sub": r[idx["子任务"]].strip(),
                "sql": sql,
            }
        )
    return out


def load_final_hot_sql():
    text = FINAL_SQL_PATH.read_text(encoding="utf-8")
    parts = re.split(r"^-- (黑域名|黑文件|黑 URL)\s*$", text, flags=re.M)
    mapping = {"黑域名": "black_domain_stat", "黑文件": "blcak_file_stat", "黑 URL": "blcak_url_stat"}
    final = {}
    for i in range(1, len(parts), 2):
        body = [ln for ln in parts[i + 1].strip().splitlines()]
        while body and body[0].startswith("--"):
            body.pop(0)
        final[mapping[parts[i]]] = "\n".join(body).strip()
    missing = set(mapping.values()) - set(final)
    if missing:
        sys.exit(f"{FINAL_SQL_PATH} 缺少语句: {sorted(missing)}")
    return final


def style_sheet(ws, widths):
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def color_status(ws, col):
    for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
        cell = row[0]
        cell.fill = DONE_FILL if cell.value == DONE else TODO_FILL


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    if OUT_PATH.exists() and not args.force:
        sys.exit(f"{OUT_PATH} 已存在，未覆盖。需要重建请加 --force。")

    units = load_initial_rows()
    final_hot = load_final_hot_sql()

    wb = Workbook()
    ws_note = wb.active
    ws_note.title = "说明"
    notes = [
        ["项目", "内容"],
        ["用途", "按业务记录每条任务的初始 SQL（旧平台 Hive/Spark SQL）和最终跑通的 StarRocks SQL。"],
        ["业务口径", "「业务分类」列取 917 梳理 CSV 的同名列，共 %d 个业务。CSV 中该列为空的行不在本表内。" % len({u["biz"] for u in units})],
        ["初始 SQL 来源", "大数据平台YARN上的任务梳理-云脑大数据平台-917梳理.csv「语句」列。该 CSV 首列带 * 的任务有多条 SQL，只举例一条；语句列为空或 / 的行，初始 SQL 留空。"],
        ["最终 SQL 来源", "目前只有热门威胁已记录，来自 outputs/topthreat-hot-stat-starrocks.sql（位于 cursor/sr-convert-all-sql-02a9 分支）。"],
        ["状态含义", "已跑通：用户已确认。未记录：尚未登记最终 SQL，不表示一定没有跑通，也不表示跑不通。"],
        ["证据边界", "已跑通不等于已切流。最终 SQL 里的日期和小时是实验窗口，不是生产调度参数。表和字段名按 docs/StarRocks语句转换规则.md 已确认的口径。未执行生产查询、切流或停服。"],
        ["补录方法", "在「SQL明细」把最终 SQL、状态、跑通日期、环境与证据填进对应行，再同步「业务总览」状态。重新运行 scripts/build_sql_run_record.py 默认拒绝覆盖本文件。"],
    ]
    for row in notes:
        ws_note.append(row)
    style_sheet(ws_note, [16, 120])
    ws_note.auto_filter.ref = None

    ws = wb.create_sheet("SQL明细")
    ws.append(
        ["业务", "业务类别(CSV)", "服务名称", "子任务", "917任务类型", "负责人",
         "初始SQL（旧平台）", "最终跑通SQL（StarRocks）", "状态", "跑通日期", "跑通环境与证据", "改写要点"]
    )
    for u in units:
        is_hot = u["biz"] == HOT
        final = clip(final_hot[u["sub"]]) if is_hot else ""
        ws.append(
            [
                u["biz"], u["category"], u["service"], u["sub"], u["status"], u["owner"],
                clip(u["sql"]) if u["sql"] else "",
                final,
                DONE if is_hot else TODO,
                "2026-10-16" if is_hot else "",
                HOT_EVIDENCE if is_hot else "",
                HOT_NOTES[u["sub"]] if is_hot else "",
            ]
        )
    style_sheet(ws, [22, 14, 28, 22, 12, 10, 70, 70, 10, 12, 42, 56])
    color_status(ws, 9)

    groups = OrderedDict()
    for u in units:
        groups.setdefault(u["biz"], []).append(u)
    ws2 = wb.create_sheet("业务总览", 1)
    ws2.append(["业务", "业务类别(CSV)", "服务", "子任务数", "有初始SQL的子任务数", "已跑通子任务数", "状态", "备注"])
    ordered = sorted(groups, key=lambda b: (b != HOT, b))
    for biz in ordered:
        g = groups[biz]
        done = len(g) if biz == HOT else 0
        status = DONE if biz == HOT else TODO
        ws2.append(
            [
                biz,
                "、".join(dict.fromkeys(x["category"] for x in g if x["category"])),
                "、".join(dict.fromkeys(x["service"] for x in g)),
                len(g),
                sum(1 for x in g if x["sql"]),
                done,
                status,
                "三条语句（黑域名、黑文件、黑 URL）均已跑通" if biz == HOT else "",
            ]
        )
    style_sheet(ws2, [28, 18, 40, 10, 14, 12, 10, 40])
    color_status(ws2, 7)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"businesses={len(groups)} units={len(units)} -> {OUT_PATH}")


if __name__ == "__main__":
    main()
