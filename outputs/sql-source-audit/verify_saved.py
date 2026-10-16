"""Verify exported values and render representative saved cells after native renderer failure."""
import json,math
from pathlib import Path
from openpyxl import load_workbook
from PIL import Image,ImageDraw,ImageFont
root=Path(__file__).resolve().parent
d=json.loads((root/'audit.json').read_text(encoding='utf-8'))
p=root/'逐任务SQL与手工来源核对.xlsx'
w=load_workbook(p,data_only=False)
assert w.sheetnames==['核对说明','逐任务核对','逐表差异','SQL片段索引','SQL原文']
assert w['逐任务核对'].max_row==46
assert {w['逐任务核对'].cell(i,1).value for i in range(2,47)}=={t['task'] for t in d['tasks']}
assert w['逐表差异'].max_row==len(d['details'])+1
assert w['SQL片段索引'].max_row==82
for i,t in enumerate(d['tasks'],2):
    s=w['逐任务核对']
    assert s.cell(i,2).value==t['conclusion']
    assert s.cell(i,3).value==t['sql_count']
    assert s.cell(i,12).value==t['manual_raw']
collected={}
for row in w['SQL原文'].iter_rows(min_row=2,values_only=True):
    collected.setdefault(row[0],[]).append(row[3] or '')
for u in d['units']:
    assert '\n'.join(collected[u['id']])==u['sql'],u['id']
assert all(w[n].freeze_panes for n in w.sheetnames)
assert all(len(w[n].tables)==1 for n in w.sheetnames)
cached=load_workbook(p,data_only=True)
assert cached['核对说明']['B26'].value==45
print('Saved workbook: 45 tasks, 81 SQL fragments, all raw SQL and original source text preserved; summary total=45.')

# Use exported cell styles for a bounded visual fallback. This is not a native Excel screenshot.
fontpath=Path('C:/Windows/Fonts/msyh.ttc')
assert fontpath.exists()
font=ImageFont.truetype(str(fontpath),14)
bold=ImageFont.truetype('C:/Windows/Fonts/msyhbd.ttc',14)
def rgb(color,default):
    if color and color.type=='rgb' and isinstance(color.rgb,str):return '#'+color.rgb[-6:]
    return default
def wrap(text,draw,f,width):
    out=[]
    for line in str(text or '').split('\n'):
        cur=''
        for ch in line:
            if cur and draw.textlength(cur+ch,font=f)>width:
                out.append(cur);cur=''
            cur+=ch
        out.append(cur)
    return out
views=[('核对说明',19,26,1,2),('逐任务核对',1,4,1,5),('逐表差异',1,4,1,3),('SQL片段索引',1,4,1,5),('SQL原文',1,8,1,4)]
for num,(name,r1,r2,c1,c2) in enumerate(views,1):
    s=w[name]
    widths=[int((s.column_dimensions[s.cell(1,c).column_letter].width or 13)*7) for c in range(c1,c2+1)]
    heights=[max(28,int((s.row_dimensions[r].height or 15)*4/3)) for r in range(r1,r2+1)]
    image=Image.new('RGB',(sum(widths)+24,sum(heights)+50),'white');draw=ImageDraw.Draw(image)
    draw.text((12,8),name,font=bold,fill='#243247')
    y=38
    for r,h in zip(range(r1,r2+1),heights):
        x=12
        for c,width in zip(range(c1,c2+1),widths):
            cell=s.cell(r,c);f=bold if cell.font.bold else font
            fill=rgb(cell.fill.fgColor,'white') if cell.fill.patternType=='solid' else 'white'
            color=rgb(cell.font.color,'#243247')
            draw.rectangle((x,y,x+width,y+h),fill=fill)
            value=cached[name].cell(r,c).value if cell.data_type=='f' else cell.value
            lines=wrap(value,draw,f,width-12)
            for j,line in enumerate(lines):draw.text((x+6,y+5+j*20),line,font=f,fill=color)
            draw.line((x,y+h,x+width,y+h),fill='#D7DEE8')
            x+=width
        y+=h
    image.save(root/f'saved-preview-{num}.png')
print('Rendered 5 representative sheet views from saved cell styles using Pillow fallback.')
