import csv, json
from pathlib import Path
from openpyxl import load_workbook
base=Path(__file__).resolve().parents[2]
for name in ['大数据平台YARN上的任务梳理-云脑大数据平台.csv','大数据平台YARN上的任务梳理-项目数据来源及输出.csv']:
    with (base/name).open(encoding='utf-8-sig',newline='') as f:
        reader=csv.reader(f)
        if '云脑' in name: next(reader)
        headers=next(reader)
        print(name, json.dumps(list(enumerate(headers)),ensure_ascii=False))
        print('sample',json.dumps(next(reader),ensure_ascii=False)[:900])
for p in base.glob('*.xlsx'):
    w=load_workbook(p,read_only=True,data_only=False)
    print('workbook',p.name,[(s.title,s.max_row,s.max_column) for s in w])
    for s in w:
        print(s.title,list(s.values)[:2][:1])
